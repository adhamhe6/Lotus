import pandas as pd
import numpy as np
import os
import re
import PyPDF2
from pathlib import Path
from deep_translator import GoogleTranslator
import nltk
from nltk.corpus import wordnet
import requests
import json
from django.conf import settings
from io import BytesIO
import openpyxl
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from openpyxl.styles import Font, PatternFill, Alignment
import warnings
warnings.filterwarnings('ignore')

# Download NLTK data
try:
    nltk.data.find('corpora/wordnet')
except LookupError:
    nltk.download('wordnet')

# Configuration
PERPLEXITY_API_KEY = getattr(settings, 'PERPLEXITY_API_KEY', 'your-api-key-here')
PERPLEXITY_API_URL = 'https://api.perplexity.ai/chat/completions'

# Translations
TRANSLATIONS = {
    'ar': {
        'title': 'أداة تحليل المحتوى',
        'keyword': 'الكلمة الرئيسية',
        'frequency': 'التكرار',
        'percentage': 'النسبة المئوية',
        'results': 'النتائج',
        'summary': 'ملخص',
        'per_file_analysis': 'التحليل حسب الملف',
    },
    'en': {
        'title': 'Content Analysis Tool',
        'keyword': 'Keyword',
        'frequency': 'Frequency',
        'percentage': 'Percentage',
        'results': 'Results',
        'summary': 'Summary',
        'per_file_analysis': 'Per-File Analysis',
    }
}

def get_text(key, lang='en'):
    """Get translated text"""
    return TRANSLATIONS.get(lang, TRANSLATIONS['en']).get(key, key)

def extract_text_from_pdf(file):
    """Extract text from PDF file"""
    try:
        text = ""
        pdf_reader = PyPDF2.PdfReader(file)
        for page in pdf_reader.pages:
            text += page.extract_text()
        return text
    except Exception as e:
        raise Exception(f"Error reading PDF: {str(e)}")

def extract_text_from_excel(file):
    """Extract text from Excel file"""
    try:
        df = pd.read_excel(file)
        return df.to_string()
    except Exception as e:
        raise Exception(f"Error reading Excel: {str(e)}")

def extract_text_from_csv(file):
    """Extract text from CSV file"""
    try:
        df = pd.read_csv(file)
        return df.to_string()
    except Exception as e:
        raise Exception(f"Error reading CSV: {str(e)}")

def extract_text_from_file(file, filename):
    """Extract text from any supported file type"""
    if filename.endswith('.pdf'):
        return extract_text_from_pdf(file)
    elif filename.endswith(('.xlsx', '.xls')):
        return extract_text_from_excel(file)
    elif filename.endswith('.csv'):
        return extract_text_from_csv(file)
    elif filename.endswith('.txt'):
        return file.read().decode('utf-8')
    else:
        raise Exception(f"Unsupported file type: {filename}")

def get_synonyms(word):
    """Get synonyms for a word using WordNet"""
    synonyms = set()
    try:
        for syn in wordnet.synsets(word.lower()):
            for lemma in syn.lemmas():
                synonyms.add(lemma.name().replace('_', ' '))
    except:
        pass
    return list(synonyms)

def analyze_content(file_contents, keywords, include_synonyms=False):
    """Analyze content for keyword frequencies"""
    results = {
        'summary': [],
        'per_file': {},
        'total_files': len(file_contents)
    }
    
    # Build search terms
    search_terms = {}
    for keyword in keywords:
        if include_synonyms:
            synonyms = get_synonyms(keyword)
            search_terms[keyword] = [keyword] + synonyms
        else:
            search_terms[keyword] = [keyword]
    
    # Analyze each file
    for file_name, content in file_contents.items():
        results['per_file'][file_name] = {}
        content_lower = content.lower()
        
        for keyword, terms in search_terms.items():
            count = 0
            for term in terms:
                pattern = r'\b' + re.escape(term.lower()) + r'\b'
                count += len(re.findall(pattern, content_lower))
            
            results['per_file'][file_name][keyword] = count
    
    # Calculate summary
    total_occurrences = {}
    for keyword in keywords:
        total = sum(results['per_file'][f].get(keyword, 0) for f in file_contents.keys())
        total_occurrences[keyword] = total
    
    # Calculate percentages
    grand_total = sum(total_occurrences.values())
    for keyword, count in total_occurrences.items():
        percentage = (count / grand_total * 100) if grand_total > 0 else 0
        results['summary'].append({
            'keyword': keyword,
            'count': count,
            'percentage': round(percentage, 2)
        })
    
    # Sort by frequency
    results['summary'].sort(key=lambda x: x['count'], reverse=True)
    
    return results

def create_pdf_report(analysis_results, language='en'):
    """Create PDF report in memory"""
    try:
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.5*inch, bottomMargin=0.5*inch)
        elements = []
        styles = getSampleStyleSheet()
        
        # Title
        title_text = get_text('title', language)
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#1f4788'),
            spaceAfter=30,
            alignment=TA_CENTER
        )
        elements.append(Paragraph(title_text, title_style))
        
        # Summary table
        summary_title = get_text('summary', language)
        elements.append(Paragraph(summary_title, styles['Heading2']))
        
        # Prepare table data
        table_data = [[
            get_text('keyword', language),
            get_text('frequency', language),
            get_text('percentage', language)
        ]]
        
        for item in analysis_results['summary']:
            keyword = item['keyword']
            if language == 'ar':
                try:
                    keyword = GoogleTranslator(source='en', target='ar').translate(keyword)
                except:
                    pass
            table_data.append([
                keyword,
                str(item['count']),
                f"{item['percentage']:.2f}%"
            ])
        
        # Create table
        summary_table = Table(table_data, colWidths=[2*inch, 2*inch, 2*inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f4788')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        elements.append(summary_table)
        
        doc.build(elements)
        buffer.seek(0)
        return buffer
    except Exception as e:
        raise Exception(f"Error creating PDF: {str(e)}")

def create_excel_report(analysis_results, language='en'):
    """Create Excel report in memory"""
    try:
        buffer = BytesIO()
        wb = openpyxl.Workbook()
        
        # Summary Sheet
        ws_summary = wb.active
        ws_summary.title = "Summary Report" if language == 'en' else "ملخص التقرير"
        
        # Headers
        headers = [
            get_text('keyword', language),
            get_text('frequency', language),
            get_text('percentage', language)
        ]
        ws_summary.append(headers)
        
        # Style headers
        header_fill = PatternFill(start_color="1f4788", end_color="1f4788", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws_summary[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
        
        # Add data
        for item in analysis_results['summary']:
            keyword = item['keyword']
            if language == 'ar':
                try:
                    keyword = GoogleTranslator(source='en', target='ar').translate(keyword)
                except:
                    pass
            ws_summary.append([
                keyword,
                item['count'],
                f"{item['percentage']:.2f}%"
            ])
        
        # Adjust column widths
        ws_summary.column_dimensions['A'].width = 20
        ws_summary.column_dimensions['B'].width = 15
        ws_summary.column_dimensions['C'].width = 15
        
        # Per-File Analysis Sheet
        if analysis_results['per_file']:
            ws_per_file = wb.create_sheet("Per-File Analysis" if language == 'en' else "التحليل حسب الملف")
            
            # First row with file names
            first_row = ["Keyword" if language == 'en' else "الكلمة الرئيسية"]
            file_names = list(analysis_results['per_file'].keys())
            first_row.extend(file_names)
            first_row.append("Total")
            ws_per_file.append(first_row)
            
            # Style header
            for cell in ws_per_file[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center')
            
            # Add data
            for item in analysis_results['summary']:
                keyword = item['keyword']
                if language == 'ar':
                    try:
                        keyword = GoogleTranslator(source='en', target='ar').translate(keyword)
                    except:
                        pass
                row = [keyword]
                for file_name in file_names:
                    row.append(analysis_results['per_file'][file_name].get(item['keyword'], 0))
                row.append(item['count'])
                ws_per_file.append(row)
            
            # Adjust column widths
            ws_per_file.column_dimensions['A'].width = 20
            for i in range(len(file_names) + 2):
                ws_per_file.column_dimensions[chr(66 + i)].width = 12
        
        wb.save(buffer)
        buffer.seek(0)
        return buffer
    except Exception as e:
        raise Exception(f"Error creating Excel: {str(e)}")

def create_csv_report(analysis_results, language='en'):
    """Create CSV report in memory"""
    try:
        buffer = BytesIO()
        
        import pandas as pd
        summary_df = pd.DataFrame(analysis_results['summary'])
        
        if language == 'ar':
            summary_df['keyword'] = summary_df['keyword'].apply(
                lambda x: GoogleTranslator(source='en', target='ar').translate(x) if x else x
            )
        
        summary_df.columns = [
            get_text('keyword', language),
            get_text('frequency', language),
            get_text('percentage', language)
        ]
        
        buffer.write(summary_df.to_csv(index=False).encode('utf-8'))
        buffer.seek(0)
        return buffer
    except Exception as e:
        raise Exception(f"Error creating CSV: {str(e)}")

def chat_with_perplexity(message, conversation_history=None):
    """Chat with Perplexity AI with better error handling"""
    if conversation_history is None:
        conversation_history = []
    
    try:
        system_message = {
            "role": "system",
            "content": "You are an expert content analysis assistant. Help users understand keyword frequency analysis results and provide insights about the data they've uploaded."
        }
        
        messages = [system_message] + conversation_history + [
            {"role": "user", "content": message}
        ]
        
        headers = {
            "Authorization": f"Bearer {PERPLEXITY_API_KEY}",
            "Content-Type": "application/json"
        }
        
        payload = {
            "model": "pplx-7b-online",
            "messages": messages,
            "temperature": 0.7,
            "max_tokens": 1024
        }
        
        response = requests.post(PERPLEXITY_API_URL, json=payload, headers=headers, timeout=30)
        
        if response.status_code == 200:
            result = response.json()
            if 'choices' in result and len(result['choices']) > 0:
                return result['choices'][0]['message']['content']
            else:
                return "Error: No response choices available from API"
        else:
            return f"API Error: {response.status_code} - {response.text}"
            
    except requests.exceptions.Timeout:
        return "Error: Request timeout - please try again"
    except requests.exceptions.ConnectionError:
        return "Error: Cannot connect to AI service - please check your connection"
    except Exception as e:
        return f"Error: {str(e)}"
    



# #updated without perplexity
# import pandas as pd
# import numpy as np
# import os
# import re
# import PyPDF2
# from pathlib import Path
# from deep_translator import GoogleTranslator
# import nltk
# from nltk.corpus import wordnet
# from io import BytesIO
# import openpyxl
# from reportlab.lib.pagesizes import letter
# from reportlab.lib import colors
# from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
# from reportlab.lib.units import inch
# from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
# from reportlab.lib.enums import TA_CENTER, TA_LEFT
# from openpyxl.styles import Font, PatternFill, Alignment
# import warnings
# warnings.filterwarnings('ignore')

# # Download NLTK data
# try:
#     nltk.data.find('corpora/wordnet')
# except LookupError:
#     nltk.download('wordnet')

# # Translations
# TRANSLATIONS = {
#     'ar': {
#         'title': 'أداة تحليل المحتوى',
#         'keyword': 'الكلمة الرئيسية',
#         'frequency': 'التكرار',
#         'percentage': 'النسبة المئوية',
#         'results': 'النتائج',
#         'summary': 'ملخص',
#         'per_file_analysis': 'التحليل حسب الملف',
#     },
#     'en': {
#         'title': 'Content Analysis Tool',
#         'keyword': 'Keyword',
#         'frequency': 'Frequency',
#         'percentage': 'Percentage',
#         'results': 'Results',
#         'summary': 'Summary',
#         'per_file_analysis': 'Per-File Analysis',
#     }
# }

# def get_text(key, lang='en'):
#     """Get translated text"""
#     return TRANSLATIONS.get(lang, TRANSLATIONS['en']).get(key, key)

# def extract_text_from_pdf(file):
#     """Extract text from PDF file"""
#     try:
#         text = ""
#         pdf_reader = PyPDF2.PdfReader(file)
#         for page in pdf_reader.pages:
#             text += page.extract_text()
#         return text
#     except Exception as e:
#         raise Exception(f"Error reading PDF: {str(e)}")

# def extract_text_from_excel(file):
#     """Extract text from Excel file"""
#     try:
#         df = pd.read_excel(file)
#         return df.to_string()
#     except Exception as e:
#         raise Exception(f"Error reading Excel: {str(e)}")

# def extract_text_from_csv(file):
#     """Extract text from CSV file"""
#     try:
#         df = pd.read_csv(file)
#         return df.to_string()
#     except Exception as e:
#         raise Exception(f"Error reading CSV: {str(e)}")

# def extract_text_from_file(file, filename):
#     """Extract text from any supported file type"""
#     if filename.endswith('.pdf'):
#         return extract_text_from_pdf(file)
#     elif filename.endswith(('.xlsx', '.xls')):
#         return extract_text_from_excel(file)
#     elif filename.endswith('.csv'):
#         return extract_text_from_csv(file)
#     elif filename.endswith('.txt'):
#         return file.read().decode('utf-8')
#     else:
#         raise Exception(f"Unsupported file type: {filename}")

# def get_synonyms(word):
#     """Get synonyms for a word using WordNet"""
#     synonyms = set()
#     try:
#         for syn in wordnet.synsets(word.lower()):
#             for lemma in syn.lemmas():
#                 synonyms.add(lemma.name().replace('_', ' '))
#     except:
#         pass
#     return list(synonyms)

# def analyze_content(file_contents, keywords, include_synonyms=False):
#     """Analyze content for keyword frequencies"""
#     results = {
#         'summary': [],
#         'per_file': {},
#         'total_files': len(file_contents)
#     }
    
#     # Build search terms
#     search_terms = {}
#     for keyword in keywords:
#         if include_synonyms:
#             synonyms = get_synonyms(keyword)
#             search_terms[keyword] = [keyword] + synonyms
#         else:
#             search_terms[keyword] = [keyword]
    
#     # Analyze each file
#     for file_name, content in file_contents.items():
#         results['per_file'][file_name] = {}
#         content_lower = content.lower()
        
#         for keyword, terms in search_terms.items():
#             count = 0
#             for term in terms:
#                 pattern = r'\b' + re.escape(term.lower()) + r'\b'
#                 count += len(re.findall(pattern, content_lower))
            
#             results['per_file'][file_name][keyword] = count
    
#     # Calculate summary
#     total_occurrences = {}
#     for keyword in keywords:
#         total = sum(results['per_file'][f].get(keyword, 0) for f in file_contents.keys())
#         total_occurrences[keyword] = total
    
#     # Calculate percentages
#     grand_total = sum(total_occurrences.values())
#     for keyword, count in total_occurrences.items():
#         percentage = (count / grand_total * 100) if grand_total > 0 else 0
#         results['summary'].append({
#             'keyword': keyword,
#             'count': count,
#             'percentage': round(percentage, 2)
#         })
    
#     # Sort by frequency
#     results['summary'].sort(key=lambda x: x['count'], reverse=True)
    
#     return results

# def create_pdf_report(analysis_results, language='en'):
#     """Create PDF report in memory"""
#     try:
#         buffer = BytesIO()
#         doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.5*inch, bottomMargin=0.5*inch)
#         elements = []
#         styles = getSampleStyleSheet()
        
#         # Title
#         title_text = get_text('title', language)
#         title_style = ParagraphStyle(
#             'CustomTitle',
#             parent=styles['Heading1'],
#             fontSize=24,
#             textColor=colors.HexColor('#1f4788'),
#             spaceAfter=30,
#             alignment=TA_CENTER
#         )
#         elements.append(Paragraph(title_text, title_style))
        
#         # Summary table
#         summary_title = get_text('summary', language)
#         elements.append(Paragraph(summary_title, styles['Heading2']))
        
#         # Prepare table data
#         table_data = [[
#             get_text('keyword', language),
#             get_text('frequency', language),
#             get_text('percentage', language)
#         ]]
        
#         for item in analysis_results['summary']:
#             keyword = item['keyword']
#             if language == 'ar':
#                 try:
#                     keyword = GoogleTranslator(source='en', target='ar').translate(keyword)
#                 except:
#                     pass
#             table_data.append([
#                 keyword,
#                 str(item['count']),
#                 f"{item['percentage']:.2f}%"
#             ])
        
#         # Create table
#         summary_table = Table(table_data, colWidths=[2*inch, 2*inch, 2*inch])
#         summary_table.setStyle(TableStyle([
#             ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f4788')),
#             ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
#             ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
#             ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
#             ('FONTSIZE', (0, 0), (-1, 0), 12),
#             ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
#             ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
#             ('GRID', (0, 0), (-1, -1), 1, colors.black)
#         ]))
#         elements.append(summary_table)
        
#         doc.build(elements)
#         buffer.seek(0)
#         return buffer
#     except Exception as e:
#         raise Exception(f"Error creating PDF: {str(e)}")

# def create_excel_report(analysis_results, language='en'):
#     """Create Excel report in memory"""
#     try:
#         buffer = BytesIO()
#         wb = openpyxl.Workbook()
        
#         # Summary Sheet
#         ws_summary = wb.active
#         ws_summary.title = "Summary Report" if language == 'en' else "ملخص التقرير"
        
#         # Headers
#         headers = [
#             get_text('keyword', language),
#             get_text('frequency', language),
#             get_text('percentage', language)
#         ]
#         ws_summary.append(headers)
        
#         # Style headers
#         header_fill = PatternFill(start_color="1f4788", end_color="1f4788", fill_type="solid")
#         header_font = Font(bold=True, color="FFFFFF")
#         for cell in ws_summary[1]:
#             cell.fill = header_fill
#             cell.font = header_font
#             cell.alignment = Alignment(horizontal='center')
        
#         # Add data
#         for item in analysis_results['summary']:
#             keyword = item['keyword']
#             if language == 'ar':
#                 try:
#                     keyword = GoogleTranslator(source='en', target='ar').translate(keyword)
#                 except:
#                     pass
#             ws_summary.append([
#                 keyword,
#                 item['count'],
#                 f"{item['percentage']:.2f}%"
#             ])
        
#         # Adjust column widths
#         ws_summary.column_dimensions['A'].width = 20
#         ws_summary.column_dimensions['B'].width = 15
#         ws_summary.column_dimensions['C'].width = 15
        
#         # Per-File Analysis Sheet
#         if analysis_results['per_file']:
#             ws_per_file = wb.create_sheet("Per-File Analysis" if language == 'en' else "التحليل حسب الملف")
            
#             # First row with file names
#             first_row = ["Keyword" if language == 'en' else "الكلمة الرئيسية"]
#             file_names = list(analysis_results['per_file'].keys())
#             first_row.extend(file_names)
#             first_row.append("Total")
#             ws_per_file.append(first_row)
            
#             # Style header
#             for cell in ws_per_file[1]:
#                 cell.fill = header_fill
#                 cell.font = header_font
#                 cell.alignment = Alignment(horizontal='center')
            
#             # Add data
#             for item in analysis_results['summary']:
#                 keyword = item['keyword']
#                 if language == 'ar':
#                     try:
#                         keyword = GoogleTranslator(source='en', target='ar').translate(keyword)
#                     except:
#                         pass
#                 row = [keyword]
#                 for file_name in file_names:
#                     row.append(analysis_results['per_file'][file_name].get(item['keyword'], 0))
#                 row.append(item['count'])
#                 ws_per_file.append(row)
            
#             # Adjust column widths
#             ws_per_file.column_dimensions['A'].width = 20
#             for i in range(len(file_names) + 2):
#                 ws_per_file.column_dimensions[chr(66 + i)].width = 12
        
#         wb.save(buffer)
#         buffer.seek(0)
#         return buffer
#     except Exception as e:
#         raise Exception(f"Error creating Excel: {str(e)}")

# def create_csv_report(analysis_results, language='en'):
#     """Create CSV report in memory"""
#     try:
#         buffer = BytesIO()
        
#         import pandas as pd
#         summary_df = pd.DataFrame(analysis_results['summary'])
        
#         if language == 'ar':
#             summary_df['keyword'] = summary_df['keyword'].apply(
#                 lambda x: GoogleTranslator(source='en', target='ar').translate(x) if x else x
#             )
        
#         summary_df.columns = [
#             get_text('keyword', language),
#             get_text('frequency', language),
#             get_text('percentage', language)
#         ]
        
#         buffer.write(summary_df.to_csv(index=False).encode('utf-8'))
#         buffer.seek(0)
#         return buffer
#     except Exception as e:
#         raise Exception(f"Error creating CSV: {str(e)}")

# def local_ai_chat(message, analysis_results, conversation_history=None):
#     """Local AI chat using analysis results - No external API needed"""
#     if conversation_history is None:
#         conversation_history = []
    
#     try:
#         # Convert message to lowercase for easier matching
#         message_lower = message.lower().strip()
        
#         # Get analysis insights
#         summary = analysis_results.get('summary', [])
#         per_file = analysis_results.get('per_file', {})
#         total_files = analysis_results.get('total_files', 0)
        
#         # Response templates
#         responses = {
#             'greeting': [
#                 "Hello! I'm your content analysis assistant. I can help you understand your keyword analysis results.",
#                 "Hi there! I'm here to help you interpret your content analysis findings.",
#                 "Welcome! I can provide insights about your keyword frequency analysis."
#             ],
#             'help': [
#                 "I can help you with:\n• Keyword frequency insights\n• File-by-file analysis\n• Trend interpretation\n• Report explanations\n\nJust ask me about your analysis!",
#                 "You can ask me about:\n• Most frequent keywords\n• Keyword distribution across files\n• Analysis methodology\n• How to interpret the results"
#             ],
#             'summary': f"I analyzed {total_files} files and found {len(summary)} keywords. ",
#             'top_keywords': "",
#             'file_analysis': "",
#             'methodology': "I use exact word matching to count keyword frequencies. For synonym analysis, I also include related words from WordNet.",
#             'unknown': "I'm not sure how to answer that. Try asking about keyword frequencies, file analysis, or interpretation of your results."
#         }
        
#         # Generate top keywords response
#         if summary:
#             top_3 = summary[:3]
#             top_keywords_text = "The most frequent keywords are: "
#             top_keywords_text += ", ".join([f"'{item['keyword']}' ({item['count']} occurrences)" for item in top_3])
#             responses['top_keywords'] = top_keywords_text
            
#             # Add percentage information
#             if len(summary) > 1:
#                 total_occurrences = sum(item['count'] for item in summary)
#                 if total_occurrences > 0:
#                     top_percentage = (top_3[0]['count'] / total_occurrences) * 100
#                     responses['top_keywords'] += f". The top keyword '{top_3[0]['keyword']}' represents {top_percentage:.1f}% of all keyword mentions."
        
#         # Generate file analysis response
#         if per_file:
#             file_with_most_keywords = max(per_file.items(), key=lambda x: sum(x[1].values()))[0]
#             responses['file_analysis'] = f"The file '{file_with_most_keywords}' contains the highest number of keyword mentions across your documents."
        
#         # Combine all information
#         full_analysis = responses['summary']
#         if responses['top_keywords']:
#             full_analysis += responses['top_keywords']
#         if responses['file_analysis']:
#             full_analysis += " " + responses['file_analysis']
        
#         # Message pattern matching
#         if any(word in message_lower for word in ['hello', 'hi', 'hey', 'greetings']):
#             return np.random.choice(responses['greeting'])
        
#         elif any(word in message_lower for word in ['help', 'what can you do', 'capabilities']):
#             return np.random.choice(responses['help'])
        
#         elif any(word in message_lower for word in ['summary', 'overview', 'results']):
#             return full_analysis
        
#         elif any(word in message_lower for word in ['top', 'most frequent', 'highest']):
#             return responses['top_keywords'] if responses['top_keywords'] else "No keyword data available for analysis."
        
#         elif any(word in message_lower for word in ['file', 'document', 'which file']):
#             return responses['file_analysis'] if responses['file_analysis'] else "No file analysis data available."
        
#         elif any(word in message_lower for word in ['how', 'method', 'methodology', 'work']):
#             return responses['methodology']
        
#         elif any(word in message_lower for word in ['keyword', 'word']):
#             # Try to find specific keyword in message
#             for item in summary:
#                 if item['keyword'].lower() in message_lower:
#                     return f"The keyword '{item['keyword']}' appears {item['count']} times ({item['percentage']}% of total)."
            
#             # If no specific keyword found, return general keyword info
#             return responses['top_keywords'] if responses['top_keywords'] else full_analysis
        
#         elif any(word in message_lower for word in ['percentage', 'percent', 'distribution']):
#             if summary:
#                 distribution = ", ".join([f"'{item['keyword']}': {item['percentage']}%" for item in summary[:5]])
#                 return f"Keyword distribution: {distribution}"
#             else:
#                 return "No percentage data available."
        
#         elif any(word in message_lower for word in ['total', 'count', 'how many']):
#             total_keywords = sum(item['count'] for item in summary)
#             return f"Total keyword occurrences across all files: {total_keywords}"
        
#         else:
#             # For unknown queries, provide the full analysis summary
#             return full_analysis + "\n\n" + responses['help'][0]
            
#     except Exception as e:
#         return f"I encountered an error while analyzing your query: {str(e)}. Please try asking about your analysis results in a different way."